"""Core checks. Run from the project root:  python -m pytest -q"""
import os
import sys
import tempfile
from pathlib import Path

os.environ["AUDITPILOT_DB"] = str(Path(tempfile.mkdtemp()) / "test.db")
os.environ["AUDITPILOT_TODAY"] = "2026-10-29"
os.environ["AI_ENABLED"] = "false"
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from core import db  # noqa: E402
from core.consolidate import norm_name  # noqa: E402
from core.generate_data import generate  # noqa: E402
from core.pipeline import resolve_conflict, run_all  # noqa: E402
from core.planner import generate_plan  # noqa: E402
from core.risk import score_supplier, what_if  # noqa: E402

STATS = run_all(regenerate=False, notify_changes=False)  # uses the shipped data/raw files
G = db.read_table("golden")


def hero(prefix):
    return G[G.legal_name.str.startswith(prefix)].iloc[0]


def test_normalizer_strips_suffixes():
    assert norm_name("Acme Pvt. Ltd.") == norm_name("ACME PRIVATE LIMITED") == "ACME"


def test_golden_records_and_dedupe():
    assert STATS["golden_records"] == 200
    assert STATS["ariba_duplicates_removed"] == 5


def test_matching_accuracy():
    truth = generate(Path(tempfile.mkdtemp()))
    m = db.read_table("matches")
    wrong = 0
    for r in m.itertuples():
        got = r.ariba_id if r.status in ("auto", "review", "approved") else None
        wrong += (truth[(r.source, r.source_id)] or None) != (got or None)
    assert wrong / len(m) < 0.02


def test_worked_example_from_plan():
    row = dict(data_access_level="pii", users_with_system_access=30, criticality_tier="high", assessment_score=70,
               open_critical=0, open_high=1, open_medium=0, last_audit_rating="Satisfactory",
               certification_status="valid", data_quality=0.9)
    res = score_supplier(row, 80, [])
    assert res["risk_score"] == 61.5 and res["risk_tier"] == "High"


def test_pii_without_cert_floor():
    row = dict(data_access_level="pii", users_with_system_access=0, criticality_tier="low", assessment_score=95,
               open_critical=0, open_high=0, open_medium=0, certification_status="none", data_quality=1.0)
    res = score_supplier(row, 5, [])
    assert res["risk_score"] == 60 and res["overrides"]


def test_hero_suppliers():
    n, o, lo = hero("Nimbus"), hero("Orion"), hero("Lotus")
    assert n.risk_tier == "High" and n.pending_conflicts == 1
    assert what_if(G, n.golden_id, "data_access_level", "internal")["risk_tier"] == "Medium"
    assert o.risk_tier == "Critical" and o.eligible and not isinstance(o.last_audit_date, str)
    assert lo.eligibility_status == "Monitor only"


def test_conflict_resolution_and_plan():
    n = hero("Nimbus")
    out = resolve_conflict(n.golden_id, "data_access_level", "pii", "Approved rule value", "Priya Raman")
    assert out["ok"] and out["after"]["risk_tier"] == "High"
    plan = generate_plan(n.golden_id, "Priya Raman", use_ai_checklist=False)
    assert plan["created"] and plan["tasks"] == 8
    assert not generate_plan(n.golden_id, "Priya Raman")["created"]  # idempotent
    notes = db.query("SELECT DISTINCT recipient FROM notifications WHERE event='audit_plan_created'")
    assert len(notes) == 3
    assert len(db.query("SELECT 1 FROM audit_log WHERE action='conflict_resolved'")) == 1


# ---------------- AI Audit Analyst (rule-based mode) ----------------
SAMPLES = Path(__file__).resolve().parents[1] / "samples"


def test_analyst_audit_report():
    from core import analyst
    from core.doc_io import read_path
    rep = analyst.analyze(read_path(SAMPLES / "Nimbus_InfoSec_Audit_Report_2026.md"), actor="test")
    n = hero("Nimbus")
    assert rep["supplier_id"] == n.golden_id and rep["document_type"] == "Audit report"
    assert rep["metrics"]["findings"] == 6 and rep["metrics"]["critical_open"] == 1
    assert rep["metrics"]["evidence_verified_pct"] == 100
    fields = {d["field"] for d in rep["discrepancies"]}
    assert {"Certification", "Uses sub-processors"} <= fields
    si = rep["score_impact"]
    assert si["projected"]["risk_score"] > si["current"]["risk_score"]


def test_analyst_soc2_and_selfassessment():
    from core import analyst
    from core.doc_io import read_path
    soc = analyst.analyze(read_path(SAMPLES / "Orion_SOC2_Type2_Summary_2026.txt"), actor="test")
    assert soc["document_type"] == "SOC 2 report" and soc["metrics"]["findings"] == 5
    lotus = analyst.analyze(read_path(SAMPLES / "Lotus_Self_Assessment_2026.txt"), actor="test")
    assert lotus["metrics"]["findings"] == 0 and lotus["risk_level"] == "Low"


def test_analyst_dataset_and_masking():
    from core import analyst, db as _db
    from core.doc_io import read_path
    rep = analyst.analyze(read_path(SAMPLES / "External_Supplier_Risk_Feed.csv"), actor="test")
    assert rep["metrics"]["critical_open"] >= 1  # sanctions hits
    assert any(d["issue"] == "Suppliers not in the supplier master" for d in rep["data_quality"])
    ctx = analyst.Context(read_path(SAMPLES / "Nimbus_InfoSec_Audit_Report_2026.md"), _db.read_table("golden"), None, "")
    assert "Nimbus" not in ctx.masked_text and "@example.com" not in ctx.masked_text


def test_accept_findings_updates_risk():
    from core import analyst
    from core.doc_io import read_path
    rep = analyst.analyze(read_path(SAMPLES / "Orion_SOC2_Type2_Summary_2026.txt"), actor="test")
    out = analyst.accept_findings(rep["report_id"], "Priya Raman")
    assert out["ok"] and out["after"]["risk_score"] >= out["before"]["risk_score"]
    assert analyst.create_tasks(rep["report_id"], "Priya Raman")["created"] == 5


def test_report_exports():
    import io
    import openpyxl
    from core import analyst, report_export
    from core.doc_io import read_path
    for name in ("Nimbus_InfoSec_Audit_Report_2026.md", "External_Supplier_Risk_Feed.csv"):
        rep = analyst.load_report(analyst.analyze(read_path(SAMPLES / name), actor="test")["report_id"])
        pdf = report_export.to_pdf(rep)
        assert pdf[:4] == b"%PDF" and len(pdf) > 5000
        wb = openpyxl.load_workbook(io.BytesIO(report_export.to_excel(rep)))
        assert {"Summary", "Findings", "Recommendations"} <= set(wb.sheetnames)
        assert wb["Findings"].max_row == len(rep["findings"]) + 1


def test_ai_provider_selection(monkeypatch):
    from core import ai
    for k in ("AI_ENABLED", "AI_PROVIDER", "OPENAI_API_KEY", "OPENAI_MODEL", "AZURE_OPENAI_ENDPOINT",
              "AZURE_OPENAI_API_KEY", "AZURE_OPENAI_DEPLOYMENT"):
        monkeypatch.delenv(k, raising=False)
    assert ai.provider() == "none"
    monkeypatch.setenv("OPENAI_API_KEY", "sk-test")
    assert ai.provider() == "openai" and ai.model() == ai.DEFAULT_OPENAI_MODEL
    monkeypatch.setenv("AZURE_OPENAI_ENDPOINT", "https://x.openai.azure.com")
    monkeypatch.setenv("AZURE_OPENAI_API_KEY", "k")
    monkeypatch.setenv("AZURE_OPENAI_DEPLOYMENT", "dep")
    assert ai.provider() == "azure" and ai.model() == "dep"
    monkeypatch.setenv("AI_PROVIDER", "openai")
    monkeypatch.setenv("OPENAI_MODEL", "gpt-4.1-mini")
    assert ai.provider() == "openai" and ai.model() == "gpt-4.1-mini"
    monkeypatch.setenv("AI_ENABLED", "false")
    assert not ai.enabled()
