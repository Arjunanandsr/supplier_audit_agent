"""Read any supported input into a common shape: text (with sections) or a table."""
from __future__ import annotations

import io
import json
import re
from dataclasses import dataclass, field

import pandas as pd

TEXT_TYPES = {".txt", ".md", ".pdf", ".docx", ".log"}
TABLE_TYPES = {".csv", ".xlsx", ".xls", ".json", ".tsv"}
SUPPORTED = sorted(TEXT_TYPES | TABLE_TYPES)
MAX_TEXT = 200_000


@dataclass
class Document:
    name: str
    kind: str                      # "text" | "table"
    text: str = ""
    table: pd.DataFrame | None = None
    sections: list[dict] = field(default_factory=list)   # [{title, start, end}]
    meta: dict = field(default_factory=dict)


def _split_sections(text: str) -> list[dict]:
    """Headings: markdown '#', 'SECTION n', numbered '1. Title', or short ALL-CAPS lines."""
    pat = re.compile(r"^(#{1,4}\s+.+|SECTION\s+\d+.*|\d{1,2}\.\s+[A-Z][^\n]{2,60}|[A-Z][A-Z0-9 &/\-]{6,70})$", re.M)
    heads = [(m.start(), m.group(0).strip("# ").strip()) for m in pat.finditer(text)]
    if not heads:
        return [{"title": "Document", "start": 0, "end": len(text)}]
    out = []
    if heads[0][0] > 0:
        out.append({"title": "Preamble", "start": 0, "end": heads[0][0]})
    for i, (pos, title) in enumerate(heads):
        end = heads[i + 1][0] if i + 1 < len(heads) else len(text)
        out.append({"title": title[:80], "start": pos, "end": end})
    return [s for s in out if text[s["start"]:s["end"]].strip()]


def from_text(name: str, text: str) -> Document:
    text = text.replace("\r\n", "\n")[:MAX_TEXT]
    return Document(name=name, kind="text", text=text, sections=_split_sections(text),
                    meta={"chars": len(text), "words": len(text.split())})


def from_table(name: str, df: pd.DataFrame) -> Document:
    df = df.copy()
    df.columns = [str(c).strip() for c in df.columns]
    return Document(name=name, kind="table", table=df, meta={"rows": len(df), "columns": len(df.columns)})


def excel_sheets(data: bytes) -> list[str]:
    return pd.ExcelFile(io.BytesIO(data)).sheet_names


def read_bytes(name: str, data: bytes, sheet: str | None = None) -> Document:
    ext = ("." + name.rsplit(".", 1)[-1].lower()) if "." in name else ""
    if ext in (".txt", ".md", ".log"):
        return from_text(name, data.decode("utf-8", errors="replace"))
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        pages = [(p.extract_text() or "") for p in reader.pages[:60]]
        doc = from_text(name, "\n\n".join(pages))
        doc.meta["pages"] = len(reader.pages)
        if not doc.text.strip():
            doc.meta["warning"] = "No extractable text (scanned PDF?). OCR is not included in the prototype."
        return doc
    if ext == ".docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = []
        for p in d.paragraphs:
            t = p.text.strip()
            if not t:
                continue
            parts.append(f"# {t}" if p.style.name.lower().startswith("heading") else t)
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return from_text(name, "\n".join(parts))
    if ext in (".csv", ".tsv"):
        sep = "\t" if ext == ".tsv" else None
        return from_table(name, pd.read_csv(io.BytesIO(data), sep=sep, engine="python", dtype=str,
                                            keep_default_na=False))
    if ext in (".xlsx", ".xls"):
        sheets = excel_sheets(data)
        if sheet is None:  # default: the largest sheet that is not a README/cover tab
            data_sheets = [x for x in sheets if x.lower() not in ("readme", "cover", "about", "notes")] or sheets
            sizes = {x: len(pd.read_excel(io.BytesIO(data), sheet_name=x, usecols=[0])) for x in data_sheets}
            sheet = max(sizes, key=sizes.get)
        doc = from_table(f"{name} [{sheet}]", pd.read_excel(io.BytesIO(data), sheet_name=sheet, dtype=str).fillna(""))
        doc.meta |= {"sheets": sheets, "sheet": sheet}
        return doc
    if ext == ".json":
        obj = json.loads(data.decode("utf-8", errors="replace"))
        if isinstance(obj, dict):
            obj = next((v for v in obj.values() if isinstance(v, list)), [obj])
        return from_table(name, pd.json_normalize(obj).fillna("").astype(str))
    raise ValueError(f"Unsupported file type '{ext}'. Supported: {', '.join(SUPPORTED)}")


def read_path(path, sheet: str | None = None) -> Document:
    from pathlib import Path
    p = Path(path)
    return read_bytes(p.name, p.read_bytes(), sheet)
